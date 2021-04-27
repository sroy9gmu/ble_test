import os
from google.cloud import firestore
from openpyxl import (
    Workbook,
    drawing,
)
from openpyxl.chart import (
    LineChart,
    Reference,
)
from copy import deepcopy

def main():
    pulse_dict = {}
    pulse_cnt = 1
    temp_dict = {}
    temp_cnt = 1
    pulse_str = 'Pulse Signal'
    temp_str = 'Body Temperature'
    index_str = 'Sample Index'
    
    file1 = open(os.getcwd() + '/CollectionNames.txt', 'r')
    Lines = file1.readlines()
    pulse_album = Lines[0].rstrip()
    temp_album = Lines[1].rstrip()

    # Query for documents
    db = firestore.Client()
    pulse_ref = db.collection(pulse_album)
    temp_ref = db.collection(temp_album)

    # Get contents from documents
    for doc in pulse_ref.stream():
        #print(u'{} => {}'.format(doc.id, doc.to_dict()))
        tmp_dict = doc.to_dict()
        mykey = tmp_dict[index_str.lower()]
        myvalue = tmp_dict[pulse_str.lower()]
        pulse_dict[mykey] = myvalue
        pulse_cnt += 1
        
    for doc in temp_ref.stream():
        #print(u'{} => {}'.format(doc.id, doc.to_dict()))
        tmp_dict = doc.to_dict()
        mykey = tmp_dict[index_str.lower()]
        myvalue = tmp_dict[temp_str.lower()]
        temp_dict[mykey] = myvalue
        temp_cnt += 1

    # Initialize excel file for output
    wb = Workbook()
    ws0 = wb.active
    ws0.title = pulse_str
    wb.create_sheet(temp_str)
    ws1 = wb[temp_str]

    # Populate header contents
    pulse_rows = [
        [index_str, pulse_str]
    ]
    temp_rows = [
        [index_str, temp_str]
    ]

    # Add remaining values for plotting graphs
    for i in range(1, len(pulse_dict) + 1):        
        pulse_rows.append([i, pulse_dict[i]])
    for i in range(1, len(temp_dict) + 1):        
        temp_rows.append([i, temp_dict[i]])

    for row in pulse_rows:
        ws0.append(row)
    for row in temp_rows:
        ws1.append(row)

    # Set the graph properties
    c0 = LineChart()
    c0.title = "Heart Rate Sketch (pulse)"
    c0.style = 13
    c0.y_axis.title = pulse_str + ' Value (0 - 1024)'
    c0.x_axis.title = index_str
    
    c1 = LineChart()
    c1.title = "Heart Rate Sketch (temperature)"
    c1.style = 13
    c1.y_axis.title = temp_str + ' Value (Fahrenheit)'
    c1.x_axis.title = index_str

    data0 = Reference(ws0, min_col=2, min_row=1, max_col=2, max_row=pulse_cnt)
    c0.add_data(data0, titles_from_data=True)
    data1 = Reference(ws1, min_col=2, min_row=1, max_col=2, max_row=temp_cnt)
    c1.add_data(data1, titles_from_data=True)

    # Style the lines
    s0 = c0.series[0]
    s0.smooth = True # Make the line smooth
    lineProp = drawing.line.LineProperties(solidFill = drawing.colors.ColorChoice(prstClr='red'))
    s0.graphicalProperties.line = lineProp
    dest0 = "A" + str(pulse_cnt + 3)
    ws0.add_chart(c0, dest0) # Plot graph

    s1 = c1.series[0]
    s1.smooth = True # Make the line smooth
    dest1 = "A" + str(temp_cnt + 3)
    ws1.add_chart(c1, dest1) # Plot graph
    
    wb.save("HeartRatePlot.xlsx")
    
if __name__ == '__main__':
    main()